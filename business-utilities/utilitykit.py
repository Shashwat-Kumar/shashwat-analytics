"""Local business utilities. No hosted API, telemetry, eval or credential required."""
from __future__ import annotations

import base64
import csv
import difflib
import io
import json
import math
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

VERSION = "1.0.0"
META = "__utility_"
OPS = {"eq", "ne", "gt", "gte", "lt", "lte", "in", "not_in"}
NUMERIC = {"gt", "gte", "lt", "lte"}


def parse_csv(text):
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff")), strict=True)
    fields = reader.fieldnames or []
    if not fields or any(not str(f).strip() for f in fields):
        raise ValueError("CSV needs non-empty column headers")
    fields = [f.strip() for f in fields]
    if len(set(fields)) != len(fields):
        raise ValueError("CSV column headers must be unique")
    reader.fieldnames = fields
    rows = []
    for n, row in enumerate(reader, 2):
        if None in row or any(v is None for v in row.values()):
            raise ValueError(f"CSV row {n} has a different number of columns")
        rows.append(row)
    return rows


def read_table(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            data = list(book.active.values)
        finally:
            book.close()
        if not data:
            raise ValueError("Workbook is empty")
        out = io.StringIO()
        csv.writer(out).writerows([["" if v is None else str(v) for v in row] for row in data])
        return parse_csv(out.getvalue())
    return parse_csv(path.read_text(encoding="utf-8-sig"))


def safe_cell(value):
    value = str(value if value is not None else "")
    # CSV opened in spreadsheets must not execute formulas supplied by an input.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value


def csv_text(rows):
    if not rows:
        return ""
    fields = list(dict.fromkeys(k for row in rows for k in row))
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows({k: safe_cell(row.get(k, "")) for k in fields} for row in rows)
    return out.getvalue()


def check_metadata(rows):
    if any(k.startswith(META) for row in rows for k in row):
        raise ValueError(f"Input columns starting with {META} are reserved")


def number(value):
    result = float(str(value).strip())
    if not math.isfinite(result):
        raise ValueError("Expected a finite number")
    return result


def clean_rows(rows, keys=(), required=(), date_columns=()):
    """Trim values, retain first key occurrence, route invalid rows for review."""
    check_metadata(rows)
    wanted = set(keys) | set(required) | set(date_columns)
    if rows and not wanted.issubset(rows[0]):
        raise ValueError("Unknown columns: " + ", ".join(sorted(wanted - rows[0].keys())))
    seen, output = set(), []
    for index, original in enumerate(rows, 2):
        row = {k: str(v).strip() for k, v in original.items()}
        problems = [f"MISSING:{k}" for k in required if not row.get(k)]
        for key in date_columns:
            if not row.get(key):
                continue
            parsed = None
            for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
                try:
                    parsed = datetime.strptime(row[key], fmt).date().isoformat()
                    break
                except ValueError:
                    pass
            if parsed is None:
                problems.append(f"INVALID_DATE:{key}")
            else:
                row[key] = parsed
        if keys and any(not row.get(k) for k in keys):
            problems.append("MISSING_DEDUPE_KEY")
        key = tuple(row.get(k, "") for k in keys) if keys else tuple(row.items())
        status = "REVIEW" if problems else "DUPLICATE" if key in seen else "KEEP"
        if status == "KEEP":
            seen.add(key)
        row.update({META + "row": index, META + "status": status,
                    META + "reason": "; ".join(problems)})
        output.append(row)
    return output


def validate_rules(rules):
    if not rules:
        raise ValueError("At least one rule is required")
    ids = set()
    for rule in rules:
        if any(not str(rule.get(k, "")).strip() for k in ("id", "field", "op", "value", "decision", "reason")):
            raise ValueError("Each rule needs id, field, op, value, decision and reason")
        if rule["id"] in ids:
            raise ValueError("Rule IDs must be unique")
        ids.add(rule["id"])
        if rule["op"] not in OPS:
            raise ValueError(f"Unsupported operator: {rule['op']}")
        if rule["op"] in NUMERIC:
            number(rule["value"])


def rule_match(actual, rule):
    op, expected = rule["op"], rule["value"]
    if op in NUMERIC:
        actual, expected = number(actual), number(expected)
        return {"gt": actual > expected, "gte": actual >= expected,
                "lt": actual < expected, "lte": actual <= expected}[op]
    actual, expected = str(actual).strip().lower(), str(expected).strip().lower()
    if op == "eq":
        return actual == expected
    if op == "ne":
        return actual != expected
    return (actual in [x.strip() for x in expected.split("|")]) == (op == "in")


def evaluate_rows(rows, rules, default="REVIEW"):
    """Ordered first-match rules. Invalid required input always yields REVIEW."""
    validate_rules(rules)
    check_metadata(rows)
    output = []
    for original in rows:
        row, errors = dict(original), []
        for rule in rules:
            value = row.get(rule["field"])
            if value is None or not str(value).strip():
                errors.append("MISSING:" + rule["field"])
            elif rule["op"] in NUMERIC:
                try:
                    number(value)
                except ValueError:
                    errors.append("INVALID_NUMBER:" + rule["field"])
        if errors:
            decision, rid, reason = "REVIEW", "INPUT_VALIDATION", "; ".join(dict.fromkeys(errors))
        else:
            matched = next((r for r in rules if rule_match(row[r["field"]], r)), None)
            decision = matched["decision"] if matched else default
            rid = matched["id"] if matched else "DEFAULT"
            reason = matched["reason"] if matched else "No rule matched; configured default applied"
        row.update({META + "decision": decision, META + "rule": rid, META + "reason": reason})
        output.append(row)
    return output


def campaigns(brief):
    """Deterministic drafting from supplied facts; no invented rates or eligibility."""
    for key in ("brand", "product", "audience", "benefit", "cta", "disclaimer", "optout"):
        if not str(brief.get(key, "")).strip():
            raise ValueError(f"Campaign brief is missing {key}")
    b = {k: str(v).strip() for k, v in brief.items()}
    limit = int(b.get("sms_limit", 160))
    if not 1 <= limit <= 10000:
        raise ValueError("SMS character limit must be between 1 and 10,000")
    banned = [x.strip().lower() for x in str(b.get("blocked_phrases", "guaranteed approval|risk-free|100% guaranteed")).split("|") if x.strip()]
    tail = f"{b['cta']} {b['disclaimer']} {b['optout']}"
    variants = [
        ("SMS", "Direct", f"{b['brand']}: Explore {b['product']}. {b['benefit']} {tail}"),
        ("SMS", "Audience", f"For {b['audience']}: {b['product']} from {b['brand']}. {b['benefit']} {tail}"),
        ("WhatsApp", "Conversational", f"Hello!\n{b['brand']} introduces {b['product']} for {b['audience']}.\n{b['benefit']}\n{tail}"),
        ("Email", "Informative", f"Subject: Explore {b['product']} with {b['brand']}\n\nHello,\n\nFor {b['audience']}, {b['product']} offers: {b['benefit']}\n\n{b['cta']}\n\n{b['disclaimer']}\n{b['optout']}")]
    result = []
    for channel, variant, message in variants:
        flags = ["BLOCKED_PHRASE:" + term for term in banned if term in message.lower()]
        if channel == "SMS" and len(message) > limit:
            flags.append("OVER_CHARACTER_LIMIT")
        result.append({"channel": channel, "variant": variant, "message": message,
                       "characters": len(message), "status": "REVISE" if flags else "DRAFT_REVIEW_REQUIRED",
                       "flags": "; ".join(flags)})
    return result


def translate_rows(rows, memory, source="en", target="hi", column="text", engine="memory"):
    check_metadata(rows)
    if engine not in {"memory", "argos"}:
        raise ValueError("Engine must be memory or argos")
    if rows and column not in rows[0]:
        raise ValueError(f"Column not found: {column}")
    lookup = {}
    for item in memory:
        if item.get("source_language") == source and item.get("target_language") == target:
            key = item["source"].strip()
            if key in lookup and lookup[key] != item["target"]:
                raise ValueError("Conflicting translation memory entry: " + key)
            lookup[key] = item["target"]
    translator = None
    if engine == "argos" and source != target:
        import argostranslate.translate
        languages = argostranslate.translate.get_installed_languages()
        src = next((x for x in languages if x.code == source), None)
        dst = next((x for x in languages if x.code == target), None)
        if not src or not dst:
            raise ValueError("Install the requested Argos language model locally first")
        translator = src.get_translation(dst)
        if translator is None:
            raise ValueError("No installed translation route for this language pair")
    output = []
    protected = re.compile(r"\{[^{}]+\}|https?://\S+|\d+(?:[.,:/-]\d+)*")
    def translate_segment(segment):
        if not segment.strip():
            return segment
        leading = re.match(r"^\s*", segment).group()
        trailing = re.search(r"\s*$", segment).group()
        return leading + translator.translate(segment.strip()).strip() + trailing
    for original in rows:
        row, text = dict(original), str(original.get(column, "")).strip()
        translated, method = "", "UNMATCHED"
        if not text:
            method = "EMPTY"
        elif source == target:
            translated, method = text, "IDENTITY"
        elif text in lookup:
            translated, method = lookup[text], "MEMORY"
        elif translator:
            # Translate around numbers, URLs and {placeholders}, retaining these verbatim.
            pieces, last = [], 0
            for match in protected.finditer(text):
                segment = text[last:match.start()]
                pieces.append(translate_segment(segment))
                pieces.append(match.group())
                last = match.end()
            segment = text[last:]
            pieces.append(translate_segment(segment))
            translated, method = "".join(pieces), "ARGOS"
        qa = sorted(protected.findall(text)) == sorted(protected.findall(translated)) if translated else False
        row.update({META + "translation": translated, META + "method": method,
                    META + "status": "DRAFT_REVIEW_REQUIRED" if translated and qa else "REVIEW_REQUIRED",
                    META + "tokens_preserved": qa})
        output.append(row)
    return output


def compare_text(before, after):
    result = []
    a, b = before.splitlines(), after.splitlines()
    for tag, i, j, k, l in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag != "equal":
            result.append({"change": tag, "old_line": i + 1 if i < j else "",
                           "new_line": k + 1 if k < l else "",
                           "before": "\n".join(a[i:j]), "after": "\n".join(b[k:l])})
    return result


def extract_fields(text):
    patterns = {
        "invoice_number": r"(?:invoice\s*(?:no\.?|number|#))\s*[:\-]?\s*([A-Za-z0-9/-]+)",
        "date": r"\b(?:date)\s*[:\-]?\s*(\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4})",
        "total": r"\b(?:grand\s+total|total\s+amount|total)\s*[:\-]?\s*(?:INR|Rs\.?|₹)?\s*([\d,]+\.\d{2}|[\d,]+)"}
    return [{"field": name, "value": m.group(1) if (m := re.search(pattern, text, re.I)) else "",
             "status": "REVIEW_REQUIRED"} for name, pattern in patterns.items()]


def ocr_image(path, language="eng", min_confidence=70):
    if not re.fullmatch(r"[a-z]{3}(?:\+[a-z]{3})*", language):
        raise ValueError("Use Tesseract language codes such as eng or eng+hin")
    run = subprocess.run(["tesseract", str(path), "stdout", "-l", language, "--psm", "6", "tsv"],
                         capture_output=True, text=True, timeout=120, check=True)
    words = []
    lines = {}
    for row in csv.DictReader(io.StringIO(run.stdout), delimiter="\t"):
        text = (row.get("text") or "").strip()
        if not text:
            continue
        confidence = float(row["conf"])
        words.append({"text": text, "confidence": confidence,
                      "status": "REVIEW" if confidence < min_confidence else "CHECKED_BY_ENGINE",
                      "left": row["left"], "top": row["top"], "width": row["width"], "height": row["height"]})
        key = tuple(row[x] for x in ("page_num", "block_num", "par_num", "line_num"))
        lines.setdefault(key, []).append(text)
    text = "\n".join(" ".join(line) for line in lines.values())
    return {"text": text, "words": words, "fields": extract_fields(text), "review_required": True}


def pdf_merge(inputs, output):
    from pypdf import PdfWriter
    writer = PdfWriter()
    for path in inputs:
        writer.append(str(path))
    with open(output, "wb") as stream:
        writer.write(stream)


def pdf_extract(input_path, output, pages):
    from pypdf import PdfReader, PdfWriter
    reader, writer = PdfReader(input_path), PdfWriter()
    if not pages or any(p < 1 or p > len(reader.pages) for p in pages):
        raise ValueError("Pages are 1-based and must exist in the input PDF")
    for page in pages:
        writer.add_page(reader.pages[page - 1])
    with open(output, "wb") as stream:
        writer.write(stream)
